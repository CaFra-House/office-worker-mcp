"""Módulo de Compliance y Seguridad Empresarial para The Office Worker.

Implementa:
- scrub_metadata: eliminación de metadatos sensibles (autor, título, historial, editores) en PDF y Office (.docx, .xlsx, .pptx).
- protect_office: cifrado robusto con contraseña de documentos Office (.docx, .xlsx, .pptx) mediante msoffcrypto (ECMA-376 agile encryption).
- verify_pdf_signature: validación criptográfica e inspección honesta de firmas digitales en PDFs (pyhanko/pypdf).
"""
from __future__ import annotations
import json
import os
from pathlib import Path
from typing import Any, Sequence

from .security import safe_out


def scrub_metadata(
    input_path: str,
    output: str,
    fields: list[str] | str | None = None,
) -> dict[str, Any]:
    """Elimina metadatos sensibles de documentos PDF y Office (.docx, .xlsx, .pptx).

    - input_path: ruta al archivo original.
    - output: ruta destino del documento sin metadatos.
    - fields: lista opcional de campos a limpiar (ej: ['author', 'title']). Si se omite o es 'all', limpia todos.

    Devuelve dict con status, path, bytes, format y lista de campos eliminados.
    """
    input_path = os.path.abspath(os.path.expanduser(str(input_path)))
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Archivo de entrada no encontrado: {input_path}")

    out_path = safe_out(output)
    ext = Path(input_path).suffix.lower()

    # Parsear fields
    if isinstance(fields, str):
        try:
            parsed_fields = json.loads(fields or "[]")
        except Exception:
            parsed_fields = [f.strip() for f in fields.split(",") if f.strip()]
    elif isinstance(fields, Sequence):
        parsed_fields = list(fields)
    else:
        parsed_fields = []

    fields_lower = {str(f).strip().lower() for f in parsed_fields if str(f).strip()}
    scrub_all = len(fields_lower) == 0 or "all" in fields_lower

    scrubbed_fields: list[str] = []

    if ext == ".pdf":
        import fitz
        doc = fitz.open(input_path)
        existing_meta = doc.metadata or {}

        standard_keys = ["author", "title", "subject", "keywords", "creator", "producer", "creationDate", "modDate", "trapped"]
        if scrub_all:
            for k, v in existing_meta.items():
                if v:
                    scrubbed_fields.append(k)
            doc.set_metadata({})
        else:
            updated_meta = dict(existing_meta)
            for k in standard_keys:
                if k.lower() in fields_lower:
                    if updated_meta.get(k):
                        scrubbed_fields.append(k)
                    updated_meta[k] = ""
            doc.set_metadata(updated_meta)

        doc.save(out_path, garbage=4, deflate=True, clean=True)
        doc.close()

    elif ext == ".docx":
        from docx import Document
        doc = Document(input_path)
        cp = doc.core_properties

        props = [
            ("author", "author"),
            ("last_modified_by", "last_modified_by"),
            ("title", "title"),
            ("subject", "subject"),
            ("keywords", "keywords"),
            ("comments", "comments"),
            ("category", "category"),
            ("content_status", "content_status"),
            ("identifier", "identifier"),
        ]

        for prop_name, field_id in props:
            if scrub_all or field_id in fields_lower:
                val = getattr(cp, prop_name, None)
                if val:
                    scrubbed_fields.append(field_id)
                setattr(cp, prop_name, "")

        if scrub_all or "revision" in fields_lower:
            cp.revision = 1
            scrubbed_fields.append("revision")

        doc.save(out_path)

    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(input_path)
        props = wb.properties

        fields_map = [
            ("creator", "author"),
            ("lastModifiedBy", "last_modified_by"),
            ("title", "title"),
            ("subject", "subject"),
            ("description", "comments"),
            ("keywords", "keywords"),
            ("category", "category"),
        ]

        for prop_attr, field_id in fields_map:
            if scrub_all or field_id in fields_lower or prop_attr.lower() in fields_lower:
                val = getattr(props, prop_attr, None)
                if val:
                    scrubbed_fields.append(field_id)
                setattr(props, prop_attr, "")

        wb.save(out_path)
        wb.close()

    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(input_path)
        cp = prs.core_properties

        props = [
            ("author", "author"),
            ("last_modified_by", "last_modified_by"),
            ("title", "title"),
            ("subject", "subject"),
            ("keywords", "keywords"),
            ("comments", "comments"),
            ("category", "category"),
        ]

        for prop_name, field_id in props:
            if scrub_all or field_id in fields_lower:
                val = getattr(cp, prop_name, None)
                if val:
                    scrubbed_fields.append(field_id)
                setattr(cp, prop_name, "")

        prs.save(out_path)

    else:
        raise ValueError(
            f"Formato no soportado para scrub_metadata: '{ext}'. Formatos soportados: .pdf, .docx, .xlsx, .pptx"
        )

    if not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        raise RuntimeError(f"Error al escribir el archivo procesado en: {out_path}")

    return {
        "status": "ok",
        "path": out_path,
        "bytes": os.path.getsize(out_path),
        "format": ext,
        "scrubbed_fields": sorted(list(set(scrubbed_fields))),
    }


def protect_office(
    input_path: str,
    output: str,
    password: str,
) -> dict[str, Any]:
    """Protege un documento Office (.docx, .xlsx, .pptx) con cifrado estándar por contraseña.

    Utiliza msoffcrypto (ECMA-376 Agile Encryption con AES). Al intentar abrir el archivo
    sin la clave provista, Microsoft Office, LibreOffice u openpyxl rechazarán la lectura.

    - input_path: ruta al archivo office a proteger.
    - output: ruta donde se guardará el archivo protegido.
    - password: clave de apertura.

    Devuelve dict con status, path, bytes, format y encrypted=True.
    """
    if not password or not str(password).strip():
        raise ValueError("Se requiere una contraseña no vacía para protect_office.")

    input_path = os.path.abspath(os.path.expanduser(str(input_path)))
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Archivo de entrada no encontrado: {input_path}")

    out_path = safe_out(output)
    ext = Path(input_path).suffix.lower()

    if ext not in (".docx", ".xlsx", ".pptx", ".xlsm", ".xltx", ".dotx", ".pptm"):
        raise ValueError(
            f"Formato no soportado para protect_office: '{ext}'. Formatos soportados: .docx, .xlsx, .pptx. "
            f"Para proteger PDFs, utilice render_document(password=...) o pdf_manipulate(operation='encrypt', password=...)."
        )

    import msoffcrypto
    with open(input_path, "rb") as inf:
        office_file = msoffcrypto.OfficeFile(inf)
        with open(out_path, "wb") as outf:
            office_file.encrypt(str(password), outf)

    if not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        raise RuntimeError(f"Error al escribir el archivo cifrado en: {out_path}")

    return {
        "status": "ok",
        "path": out_path,
        "bytes": os.path.getsize(out_path),
        "format": ext,
        "encrypted": True,
    }


def verify_pdf_signature(input_path: str) -> dict[str, Any]:
    """Verifica firmas digitales criptográficas en un archivo PDF mediante pyhanko y pypdf.

    Inspecciona si el documento contiene campos de firma, si la integridad criptográfica
    (digest) está intacta, extrae el firmante (Common Name del certificado X.509), fecha,
    motivo y ubicación, reportando advertencias honestas sobre cadenas de confianza.

    - input_path: ruta al archivo PDF a auditar.

    Devuelve dict con has_signature, valid, intact, signatures_count, signer, date, reason, location y warnings.
    """
    input_path = os.path.abspath(os.path.expanduser(str(input_path)))
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Archivo PDF no encontrado: {input_path}")

    has_signature = False
    is_valid = False
    is_intact = False
    signer_name = None
    sig_date = None
    reason = None
    location = None
    warnings: list[str] = []
    signatures_count = 0

    try:
        import logging
        for log_name in ("pyhanko", "pyhanko_certvalidator", "asn1crypto"):
            l = logging.getLogger(log_name)
            l.setLevel(logging.CRITICAL)
            l.propagate = False
        from pyhanko.pdf_utils.reader import PdfFileReader
        from pyhanko.sign.validation import validate_pdf_signature
        from .security import run_in_worker_thread_if_async

        with open(input_path, "rb") as f:
            reader = PdfFileReader(f)
            sigs = getattr(reader, "embedded_signatures", [])
            signatures_count = len(sigs)

            if signatures_count > 0:
                has_signature = True
                sig = sigs[0]
                obj = getattr(sig, "sig_object", {}) or {}
                reason = str(obj.get("/Reason") or "") or None
                location = str(obj.get("/Location") or "") or None
                raw_date = str(obj.get("/M") or "")
                sig_date = raw_date if raw_date else None

                try:
                    status = run_in_worker_thread_if_async(validate_pdf_signature, sig)
                    is_intact = bool(status.intact)
                    is_valid = bool(status.valid)
                    if status.signing_cert and status.signing_cert.subject:
                        signer_name = status.signing_cert.subject.human_friendly

                    if not is_intact:
                        warnings.append("Signature digest mismatch: document may have been modified after signing.")
                    if status.summary() and "UNTRUSTED" in status.summary():
                        warnings.append("Certificate is self-signed or not anchored in a trusted system certificate authority.")
                except Exception as ve:
                    warnings.append(f"Validation inspection note: {ve}")
            else:
                has_signature = False
                is_valid = False
                warnings.append("Document has no digital signatures.")

    except Exception as e:
        # Fallback de inspección vía pypdf
        try:
            import pypdf
            r = pypdf.PdfReader(input_path)
            fields = r.get_fields() or {}
            sig_fields = [k for k, v in fields.items() if v.get("/FT") == "/Sig"]
            if sig_fields:
                has_signature = True
                signatures_count = len(sig_fields)
                warnings.append(f"Detected {len(sig_fields)} signature field(s), but pyhanko parser encountered: {e}")
            else:
                has_signature = False
                warnings.append("Document has no digital signatures.")
        except Exception as e2:
            warnings.append(f"Failed to inspect PDF signatures: {e2}")

    return {
        "status": "ok",
        "path": input_path,
        "has_signature": has_signature,
        "valid": is_valid,
        "intact": is_intact,
        "signatures_count": signatures_count,
        "signer": signer_name,
        "date": sig_date,
        "reason": reason,
        "location": location,
        "warnings": warnings,
    }
